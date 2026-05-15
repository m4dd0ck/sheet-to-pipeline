"""Load the workbook's raw tabs and any monthly CSV drops into DuckDB, as the source of truth.

Only the six exported columns are read from each raw tab; the formula columns are the old
process and are rebuilt in dbt. Every row keeps where it came from (tab or file, row number) so
reconciliation can point at the exact lines behind a difference.
"""

import csv
import re
import tempfile
from collections.abc import Sequence
from datetime import date, datetime
from pathlib import Path

import duckdb
from openpyxl import load_workbook

from sheet_to_pipeline.legacy import RAW_HEADER

RAW_TAB = re.compile(r"^Raw - (\d{4}-\d{2})$")
CSV_NAME = re.compile(r"^(\d{4}-\d{2})\.csv$")
SalesRow = tuple[str, str, int, date, str, str, str, int, str]


class ExtractError(ValueError):
    """Raised when an input does not look like the export it should be."""


def extract(workbook: Path, incoming: Path | None, db_path: Path) -> dict[str, int]:
    """Replace the raw schema with the workbook and any CSVs in ``incoming``.

    Returns:
        Row counts per raw table.

    Raises:
        ExtractError: On an unexpected header, or a month present in two sources.
    """
    sales, products, reps = _read_workbook(workbook)
    if incoming and incoming.is_dir():
        seen = {row[0] for row in sales}
        for path in sorted(incoming.glob("*.csv")):
            rows = _read_csv(path)
            if rows and rows[0][0] in seen:
                raise ExtractError(f"{rows[0][0]} is in both the workbook and {path.name}")
            sales += rows

    db_path.parent.mkdir(parents=True, exist_ok=True)
    with duckdb.connect(str(db_path)) as connection:
        connection.execute("create schema if not exists raw")
        connection.execute(
            """create or replace table raw.sales_lines (
                month varchar, source varchar, source_row integer, sale_date date,
                invoice varchar, customer varchar, product_code varchar, qty integer,
                rep_code varchar)"""
        )
        _bulk_insert(connection, "raw.sales_lines", sales)
        connection.execute(
            """create or replace table raw.products (
                code varchar, name varchar, category varchar, unit_price decimal(10, 2),
                effective_from date, source_row integer)"""
        )
        _bulk_insert(connection, "raw.products", products)
        connection.execute(
            "create or replace table raw.reps (rep_code varchar, name varchar, region varchar)"
        )
        _bulk_insert(connection, "raw.reps", reps)
    return {"sales_lines": len(sales), "products": len(products), "reps": len(reps)}


def _read_workbook(
    path: Path,
) -> tuple[list[SalesRow], list[tuple[object, ...]], list[tuple[object, ...]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sales: list[SalesRow] = []
        for sheet in workbook.worksheets:
            match = RAW_TAB.match(sheet.title)
            if not match:
                continue
            rows = list(sheet.iter_rows(values_only=True, max_col=len(RAW_HEADER)))
            _check_header(list(rows[0]), sheet.title)
            for number, row in enumerate(rows[1:], start=2):
                if any(cell is not None for cell in row):
                    sales.append(_sales_row(match.group(1), sheet.title, number, list(row)))
        product_rows = list(workbook["Products"].iter_rows(min_row=2, values_only=True))
        products: list[tuple[object, ...]] = [
            (*row[:4], _to_date(row[4]), n) for n, row in enumerate(product_rows, 2)
        ]
        reps: list[tuple[object, ...]] = [
            tuple(row[:3]) for row in workbook["Reps"].iter_rows(min_row=2, values_only=True)
        ]
    finally:
        workbook.close()
    return sales, products, reps


def _read_csv(path: Path) -> list[SalesRow]:
    match = CSV_NAME.match(path.name)
    if not match:
        raise ExtractError(f"Monthly drops must be named YYYY-MM.csv, got {path.name}")
    with path.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    _check_header(rows[0], path.name)
    return [
        _sales_row(match.group(1), path.name, number, list(row))
        for number, row in enumerate(rows[1:], start=2)
        if any(row)
    ]


def _check_header(header: Sequence[object], source: str) -> None:
    if [str(cell).strip() if cell else "" for cell in header[: len(RAW_HEADER)]] != RAW_HEADER:
        raise ExtractError(f"{source}: expected columns {RAW_HEADER}, got {header}")


def _sales_row(month: str, source: str, number: int, row: list[object]) -> SalesRow:
    day, invoice, customer, product, qty, rep = row[:6]
    return (
        month, source, number, _to_date(day), str(invoice).strip(), str(customer).strip(),
        str(product).strip(), int(str(qty)), str(rep).strip(),
    )  # fmt: skip


def _to_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value).strip()[:10])


def _bulk_insert(
    connection: duckdb.DuckDBPyConnection, table: str, rows: Sequence[Sequence[object]]
) -> None:
    # Reason: executemany costs milliseconds per row; a CSV round trip loads thousands at once.
    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "rows.csv"
        with path.open("w", newline="", encoding="utf-8") as handle:
            csv.writer(handle).writerows(rows)
        connection.execute(
            f"insert into {table} select * from read_csv(?, header = false, all_varchar = true)",
            [str(path)],
        )
