"""Compare what the workbook reported with what the pipeline computes, and explain the gaps.

The workbook's revenue is every pasted line at VLOOKUP's first-match price. The pipeline's is
every real line at the price in force on the sale date. So the gap splits exactly into:

    pipeline - workbook = -(repeated lines at the old price) + (real lines x price difference)

which is what ``explain`` measures, line by line, so each cause comes with the rows behind it.
"""

from dataclasses import dataclass, field
from pathlib import Path

import duckdb
from openpyxl import load_workbook

from sheet_to_pipeline.legacy import SUMMARY_HEADER

MONEY_MEASURES = {"Revenue", "Avg Invoice", "North", "South", "Central"}
TOLERANCE = 0.005


class ReconcileError(ValueError):
    """Raised when the workbook's Summary tab cannot be read."""


@dataclass
class Difference:
    measure: str
    workbook: float | str
    pipeline: float | str

    @property
    def delta(self) -> float | None:
        if isinstance(self.workbook, str) or isinstance(self.pipeline, str):
            return None
        return round(self.pipeline - self.workbook, 2)


@dataclass
class Cause:
    kind: str  # "repeated_paste" or "stale_price"
    revenue_effect: float  # pipeline minus workbook, in money
    summary: str
    evidence: list[str] = field(default_factory=list)


@dataclass
class MonthCheck:
    month: str
    label: str
    differences: list[Difference] = field(default_factory=list)
    causes: list[Cause] = field(default_factory=list)

    @property
    def revenue_gap(self) -> float:
        revenue = next((d for d in self.differences if d.measure == "Revenue"), None)
        return (revenue.delta or 0.0) if revenue else 0.0

    @property
    def unexplained(self) -> float:
        return round(self.revenue_gap - sum(c.revenue_effect for c in self.causes), 2)

    @property
    def status(self) -> str:
        if not self.differences:
            return "match"
        return "explained" if abs(self.unexplained) < 0.01 else "unexplained"


def read_workbook_summary(workbook: Path) -> dict[str, dict[str, float | str]]:
    """Summary tab rows keyed by month label ("Mar 2024")."""
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        rows = list(book["Summary"].iter_rows(values_only=True))
    finally:
        book.close()
    header_index = next(
        (i for i, row in enumerate(rows) if list(row[: len(SUMMARY_HEADER)]) == SUMMARY_HEADER),
        None,
    )
    if header_index is None:
        raise ReconcileError(f"No Summary header {SUMMARY_HEADER} found")
    summary: dict[str, dict[str, float | str]] = {}
    for row in rows[header_index + 1 :]:
        if row and row[0]:
            values = dict(zip(SUMMARY_HEADER, row, strict=False))
            summary[str(row[0])] = {
                k: v for k, v in values.items() if isinstance(v, (int, float, str))
            }
    return summary


def pipeline_summary(db_path: Path) -> dict[str, tuple[str, dict[str, float | str]]]:
    """Mart rows keyed by month label, with the YYYY-MM month alongside."""
    with duckdb.connect(str(db_path), read_only=True) as connection:
        rows = connection.execute(
            "select month, month_label, revenue, units, invoices, avg_invoice, north, south, "
            "central, top_rep from marts.mart_monthly_summary order by month"
        ).fetchall()
    result = {}
    for month, label, *values in rows:
        numbers = [float(v) if not isinstance(v, str) else v for v in values]
        result[label] = (month, dict(zip(SUMMARY_HEADER[1:], numbers, strict=True)))
    return result


def compare(workbook: Path, db_path: Path) -> list[MonthCheck]:
    """Every Summary cell that differs between the workbook and the pipeline, per month."""
    reported = read_workbook_summary(workbook)
    checks = []
    for label, (month, computed) in pipeline_summary(db_path).items():
        check = MonthCheck(month=month, label=label)
        old = reported.get(label, {})
        for measure in SUMMARY_HEADER[1:]:
            was, now = old.get(measure, "missing"), computed[measure]
            if isinstance(was, str) or isinstance(now, str):
                if was != now:
                    check.differences.append(Difference(measure, was, now))
            elif abs(now - was) > TOLERANCE:
                check.differences.append(Difference(measure, float(was), float(now)))
        checks.append(check)
    return checks


# The workbook priced every line at the first row for its code in Products (VLOOKUP's rule).
LEGACY_PRICES = """
    select product_code, unit_price as legacy_price
    from staging.stg_products
    qualify row_number() over (partition by product_code order by effective_from) = 1
"""


def explain(checks: list[MonthCheck], db_path: Path) -> list[MonthCheck]:
    """Attach the causes behind each month's revenue gap, with the rows that prove them."""
    with duckdb.connect(str(db_path), read_only=True) as connection:
        for check in checks:
            if check.differences:
                check.causes = _repeats(connection, check.month) + _stale(connection, check.month)
    return checks


def _repeats(connection: duckdb.DuckDBPyConnection, month: str) -> list[Cause]:
    row = connection.execute(
        f"""
        with legacy as ({LEGACY_PRICES})
        select count(*), sum(lines.qty * legacy.legacy_price), min(source_row),
               max(source_row), any_value(source)
        from staging.stg_sales_lines as lines
        inner join legacy using (product_code)
        where lines.month = ? and lines.is_repeat_paste
        """,
        [month],
    ).fetchone()
    if not row or not row[0]:
        return []
    count, revenue, first_row, last_row, source = row
    return [
        Cause(
            kind="repeated_paste",
            revenue_effect=-round(float(revenue), 2),
            summary=f"{count} lines were pasted twice, adding {float(revenue):,.2f} of revenue "
            "that never happened.",
            evidence=[f"{source}, rows {first_row}-{last_row} repeat earlier lines exactly"],
        )
    ]


def _stale(connection: duckdb.DuckDBPyConnection, month: str) -> list[Cause]:
    rows = connection.execute(
        f"""
        with legacy as ({LEGACY_PRICES})
        select lines.product_code, lines.product_name, legacy.legacy_price, lines.unit_price,
               sum(lines.qty), sum(lines.qty * (lines.unit_price - legacy.legacy_price))
        from marts.fct_sales_lines as lines
        inner join legacy using (product_code)
        where lines.month = ? and lines.unit_price != legacy.legacy_price
        group by all
        order by 1
        """,
        [month],
    ).fetchall()
    if not rows:
        return []
    effect = round(sum(float(r[5]) for r in rows), 2)
    return [
        Cause(
            kind="stale_price",
            revenue_effect=effect,
            summary=f"{len(rows)} products were priced at their old price; the new prices "
            f"from Products were never used, understating revenue by {effect:,.2f}.",
            evidence=[
                f"{code} {name}: {int(qty)} sold at {float(old):.2f} instead of {float(new):.2f}"
                for code, name, old, new, qty, _ in rows
            ],
        )
    ]
