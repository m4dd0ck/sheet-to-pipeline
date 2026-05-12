"""The "before": the monthly sales workbook a small wholesale bakery supplier runs on.

Each month someone pastes the order system's export into a new "Raw - YYYY-MM" tab, drags the
lookup formulas down, and types the month's totals into the Summary tab. Two things have gone
wrong, the way they do in real workbooks:

- March's export was pasted twice in part (``DUPLICATED_ROWS`` lines appear again at the end).
- Three prices changed on 2024-04-01. The new rows were added at the bottom of Products, but
  VLOOKUP returns the first match, so the workbook kept using the old prices.
"""

import random
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

BUSINESS = "Hartwell Bakery Supply"
MONTHS = [f"2024-{m:02d}" for m in range(1, 7)]
PRICE_CHANGE = date(2024, 4, 1)
DUPLICATED_MONTH = "2024-03"
DUPLICATED_ROWS = 25
RAW_HEADER = ["Date", "Invoice #", "Customer", "Product Code", "Qty", "Rep Code"]

PRODUCTS = [
    ("BR-01", "Sourdough loaf", "Bread", 4.20), ("BR-02", "Seeded rye", "Bread", 4.60),
    ("BR-03", "Baguette", "Bread", 2.10), ("BR-04", "Brioche loaf", "Bread", 5.40),
    ("BR-05", "Focaccia tray", "Bread", 9.80), ("PA-01", "Butter croissant", "Pastry", 1.35),
    ("PA-02", "Pain au chocolat", "Pastry", 1.55), ("PA-03", "Almond croissant", "Pastry", 1.95),
    ("PA-04", "Cinnamon bun", "Pastry", 1.80), ("PA-05", "Fruit danish", "Pastry", 1.70),
    ("CK-01", "Carrot cake (12 slices)", "Cake", 21.00), ("CK-02", "Lemon drizzle", "Cake", 16.50),
    ("CK-03", "Brownie tray", "Cake", 18.00), ("CK-04", "Banana bread", "Cake", 12.00),
    ("SV-01", "Cheese scone", "Savoury", 1.40), ("SV-02", "Sausage roll", "Savoury", 1.90),
    ("SV-03", "Quiche (8 slices)", "Savoury", 14.00), ("SV-04", "Spinach pasty", "Savoury", 2.60),
]  # fmt: skip
# New prices from 2024-04-01, appended below the originals like a person would.
PRICE_CHANGES = {"BR-01": 4.60, "PA-01": 1.50, "CK-01": 23.50}
REPS = [
    ("R01", "Priya Shah", "North"), ("R02", "Tom Ellis", "North"), ("R03", "Maya Osei", "South"),
    ("R04", "Luca Bianchi", "Central"), ("R05", "Hannah Reid", "Central"),
]  # fmt: skip
CAFES = [
    "Bean There", "The Daily Grind", "Crumb & Co", "Morning Glory", "Hilltop Deli", "Kettle Black",
    "Two Spoons", "Mill Lane Cafe", "The Corner Cup", "Rise Coffee", "Oak & Ember", "Blue Door",
    "Station Kiosk", "Harbour View", "Greenhouse Cafe", "Little Loaf", "The Pantry", "Steam Room",
    "Bramble", "Copper Pot", "Honeycomb", "Old Mill", "Northside Bakes", "Parkside", "Wren Cafe",
    "Latte Da", "The Hive", "Canal Coffee", "Maple Room", "Sunny Side", "Brick Lane Deli",
    "Pebble", "Riverbank", "Top Shelf", "Grain Store", "The Nook", "Foxglove", "Juniper",
    "Tea Leaf", "Market Hall",
]  # fmt: skip


@dataclass(frozen=True)
class Line:
    """One invoice line as the order system exports it."""

    day: date
    invoice: str
    customer: str
    product: str
    qty: int
    rep: str

    def cells(self) -> list[object]:
        return [self.day, self.invoice, self.customer, self.product, self.qty, self.rep]


def generate_lines(month: str, rng: random.Random) -> list[Line]:
    """A month of invoice lines: ~70 invoices of 1-4 products, reps fixed per customer."""
    year, number = (int(p) for p in month.split("-"))
    first = date(year, number, 1)
    days = ((first.replace(day=28) + timedelta(days=4)).replace(day=1) - first).days
    rep_of = {cafe: REPS[i % len(REPS)][0] for i, cafe in enumerate(CAFES)}
    lines = []
    for index in range(rng.randint(64, 78)):
        day = first + timedelta(days=rng.randrange(days))
        customer = rng.choice(CAFES)
        invoice = f"INV-{year % 100}{number:02d}-{index + 1:03d}"
        for product in rng.sample(PRODUCTS, rng.randint(1, 4)):
            qty = rng.randint(1, 3) if product[2] == "Cake" else rng.randint(6, 48)
            lines.append(Line(day, invoice, customer, product[0], qty, rep_of[customer]))
    return sorted(lines, key=lambda line: (line.day, line.invoice))


def month_lines(seed: int = 12) -> dict[str, list[Line]]:
    """Every month's export as the order system produced it (no duplicates)."""
    rng = random.Random(seed)
    return {month: generate_lines(month, rng) for month in MONTHS}


def build_legacy_workbook(path: Path, seed: int = 12) -> Path:
    """Write the workbook exactly as the business keeps it."""
    workbook = Workbook()
    products = workbook.active
    assert products is not None
    products.title = "Products"
    products.append(["Code", "Name", "Category", "Unit Price", "Effective From"])
    for code, name, category, price in PRODUCTS:
        products.append([code, name, category, price, date(2023, 1, 1)])
    for code, price in PRICE_CHANGES.items():
        name, category = next((p[1], p[2]) for p in PRODUCTS if p[0] == code)
        products.append([code, name, category, price, PRICE_CHANGE])

    reps = workbook.create_sheet("Reps")
    reps.append(["Rep Code", "Name", "Region"])
    for rep in REPS:
        reps.append(list(rep))

    for month, lines in month_lines(seed).items():
        sheet = workbook.create_sheet(f"Raw - {month}")
        sheet.append(RAW_HEADER)
        pasted = lines + (lines[:DUPLICATED_ROWS] if month == DUPLICATED_MONTH else [])
        for row, line in enumerate(pasted, start=2):
            sheet.append(line.cells())
            sheet.cell(row, 1).number_format = "dd/mm/yyyy"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
