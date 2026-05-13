from pathlib import Path

from openpyxl import load_workbook

from sheet_to_pipeline.legacy import (
    DUPLICATED_MONTH,
    DUPLICATED_ROWS,
    MONTHS,
    PRICE_CHANGES,
    build_legacy_workbook,
    legacy_summary,
    month_lines,
)


def test_workbook_has_the_expected_tabs_and_formulas(tmp_path: Path) -> None:
    workbook = load_workbook(build_legacy_workbook(tmp_path / "report.xlsx"))
    assert workbook.sheetnames == ["Summary", "Products", "Reps", *(f"Raw - {m}" for m in MONTHS)]
    raw = workbook[f"Raw - {DUPLICATED_MONTH}"]
    assert raw["G2"].value == "=VLOOKUP(D2,Products!$A:$D,4,FALSE)"
    expected_rows = len(month_lines()[DUPLICATED_MONTH]) + DUPLICATED_ROWS
    assert raw.max_row - 1 == expected_rows


def test_new_prices_sit_below_the_old_ones(tmp_path: Path) -> None:
    products = load_workbook(build_legacy_workbook(tmp_path / "report.xlsx"))["Products"]
    codes = [row[0] for row in products.iter_rows(min_row=2, values_only=True)]
    for code in PRICE_CHANGES:
        assert codes.index(code) < len(codes) - 1 - codes[::-1].index(code)


def test_summary_is_pasted_values_for_every_month() -> None:
    rows = legacy_summary()
    assert [row[0] for row in rows] == [
        "Jan 2024",
        "Feb 2024",
        "Mar 2024",
        "Apr 2024",
        "May 2024",
        "Jun 2024",
    ]
    assert all(isinstance(row[1], float) for row in rows)
