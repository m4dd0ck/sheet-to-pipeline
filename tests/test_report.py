from pathlib import Path

from openpyxl import load_workbook

from sheet_to_pipeline.legacy import SUMMARY_HEADER
from sheet_to_pipeline.reconcile import compare, explain, pipeline_summary
from sheet_to_pipeline.report import write_excel_report


def test_report_keeps_the_summary_layout_with_pipeline_values(
    built: tuple[Path, Path], tmp_path: Path
) -> None:
    workbook, db = built
    path = write_excel_report(db, explain(compare(workbook, db), db), tmp_path / "report.xlsx")
    book = load_workbook(path)
    summary = book["Summary"]
    assert [c.value for c in summary[4]] == SUMMARY_HEADER
    march = next(
        row for row in summary.iter_rows(min_row=5, values_only=True) if row[0] == "Mar 2024"
    )
    assert march[1] == pipeline_summary(db)["Mar 2024"][1]["Revenue"]
    statuses = [row[1] for row in book["Checks"].iter_rows(min_row=2, values_only=True)]
    assert statuses == ["match", "match", "explained", "explained", "explained", "explained"]
